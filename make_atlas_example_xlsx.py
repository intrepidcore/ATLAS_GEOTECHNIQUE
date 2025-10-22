#!/usr/bin/env python3
"""
Générateur de fichier Excel d'exemple pour l'import bulk Atlas
Format: multi-feuilles avec granulométrie "large" (profondeurs en colonnes)
"""

from openpyxl import Workbook

def add_sheet(wb, name, headers, rows):
    """Ajoute une feuille avec entêtes et données"""
    ws = wb.create_sheet(title=name)
    for col_idx, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=h)
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)

def main():
    wb = Workbook()
    wb.remove(wb.active)  # retire la feuille vide par défaut

    # 1) sondages (sites)
    add_sheet(
        wb, "sondages",
        ["code_site","localite","date","lat","lon","adm1","adm2","adm3","source"],
        [
            ["Sanfatoute","Sanfatoute","2025-10-20",None,None,None,None,None,"Campagne 2025"],
            ["Korbongou","Korbongou","2025-10-20",None,None,None,None,None,"Campagne 2025"],
        ]
    )

    # 2) echantillons (par profondeur)
    add_sheet(
        wb, "echantillons",
        ["code_site","depth_m","date","laboratory","norm","rho_s_gcm3","water_content_w","is_index","eg","commentaire"],
        [
            ["Sanfatoute",1.0,"2025-10-20","Labo Atlas",None,2.47,9.82,0.171,None,""],
            ["Sanfatoute",1.5,"2025-10-20","Labo Atlas",None,2.39,10.35,0.241,None,""],
            ["Sanfatoute",2.0,"2025-10-20","Labo Atlas",None,2.63,8.02,0.247,None,""],
            ["Korbongou",1.0,"2025-10-20","Labo Atlas",None,2.47,8.17,0.206,None,""],
            ["Korbongou",1.5,"2025-10-20","Labo Atlas",None,2.51,7.52,0.225,None,""],
            ["Korbongou",2.0,"2025-10-20","Labo Atlas",None,2.57,6.21,0.256,None,""],
        ]
    )

    # 3) atterberg (WL/WP, IP calculé en DB)
    add_sheet(
        wb, "atterberg",
        ["code_site","depth_m","wl","wp"],
        [
            ["Sanfatoute",1.0,57.57,25.06],
            ["Sanfatoute",1.5,42.95,29.96],
            ["Sanfatoute",2.0,32.46,23.83],
            ["Korbongou",1.0,39.60,16.66],
            ["Korbongou",1.5,33.47,16.09],
            ["Korbongou",2.0,24.24,16.22],
        ]
    )

    # 4) vbs
    add_sheet(
        wb, "vbs",
        ["code_site","depth_m","vbs","commentaire"],
        [
            ["Sanfatoute",1.0,7.31,"Sol argileux"],
            ["Sanfatoute",1.5,6.28,"Sol argileux"],
            ["Sanfatoute",2.0,4.69,"Sol limoneux de plasticité moyenne"],
            ["Korbongou",1.0,4.40,"Sol limoneux de plasticité moyenne"],
            ["Korbongou",1.5,3.98,"Sol limoneux de plasticité moyenne"],
            ["Korbongou",2.0,2.48,"Sol sableux argileux peu plastique"],
        ]
    )

    # 5) proctor (structure vide pour l'instant)
    add_sheet(
        wb, "proctor",
        ["code_site","depth_m","proctor_type","gamma_d_max","w_opt"],
        []
    )

    # 6) granulo_tamisage_large (passant %)
    tamis_headers = ["sieve_mm","Sanfatoute@1","Sanfatoute@1.5","Sanfatoute@2","Korbongou@1"]
    tamis_rows = [
        [25.0, 100.00, 100.00, 100.00, None],
        [20.0, 100.00, 100.00, 100.00, None],
        [16.0, 100.00, 100.00, 100.00, 100.00],
        [12.5, 99.75, 100.00,  99.47,  99.63],
        [10.0, 98.87,  99.70,  99.47,  99.33],
        [8.0,  98.50,  99.70,  99.40,  99.16],
        [6.3,  98.22,  99.61,  99.36,  98.98],
        [5.0,  97.95,  99.59,  99.34,  98.85],
        [4.0,  97.54,  99.51,  99.25,  98.77],
        [3.15, 96.56,  99.33,  99.05,  98.38],
        [2.5,  95.72,  99.07,  98.89,  97.98],
        [2.0,  95.09,  98.62,  98.71,  97.55],
        [1.6,  94.58,  98.03,  98.54,  96.88],
        [1.25, 94.21,  97.38,  98.28,  95.85],
        [1.0,  93.96,  96.87,  98.03,  94.84],
        [0.8,  93.70,  96.29,  97.75,  93.71],
        [0.63, 93.45,  95.77,  97.40,  92.59],
        [0.5,  93.06,  95.02,  96.67,  90.94],
        [0.4,  92.67,  94.46,  95.91,  89.46],
        [0.315,91.88,  93.23,  94.04,  87.40],
        [0.25, 91.00,  91.88,  91.73,  85.28],
        [0.2,  89.69,  87.92,  84.86,  82.15],
        [0.16, 88.25,  85.70,  80.61,  78.36],
        [0.125,85.50,  80.41,  73.48,  71.09],
        [0.1,  83.77,  78.65,  70.87,  68.25],
        [0.08, 82.81,  77.50,  68.75,  66.74],
    ]
    add_sheet(wb, "granulo_tamisage_large", tamis_headers, tamis_rows)

    # 7) granulo_sedimento_large (passant %)
    sed_headers = ["sieve_mm","Sanfatoute@1","Sanfatoute@1.5","Sanfatoute@2","Korbongou@1"]
    sed_rows = [
        [0.0696, 82.38, 77.12, 65.88, None],
        [0.0595, None,  None,  None,  64.51],
        [0.0494, 79.33, 73.70, 62.82, None],
        [0.0428, None,  None,  None,  61.48],
        [0.0351, 76.65, 70.28, 60.02, None],
        [0.0309, None,  None,  None,  57.69],
        [0.0224, 71.69, 65.15, 56.21, None],
        [0.0199, None,  None,  None,  53.90],
        [0.0130, 67.10, 59.33, 51.11, None],
        [0.0118, None,  None,  None,  48.59],
        [0.0076, 61.37, 53.18, 46.03, None],
        [0.0070, None,  None,  None,  42.98],
        [0.0047, 56.40, 48.05, 40.70, None],
        [0.0044, None,  None,  None,  39.51],
        [0.0033, 52.97, 43.63, 37.13, None],
        [0.0031, None,  None,  None,  37.23],
        [0.0027, 51.03, 41.23, 34.58, None],
        [0.0026, None,  None,  None,  35.41],
        [0.0024, 49.91, 39.86, 33.05, None],
        [0.0022, None,  None,  None,  33.44],
        [0.0019, 48.76, 38.50, 31.01, None],
        [0.0018, None,  None,  None,  32.39],
        [0.0014, 48.00, 37.10, 28.97, None],
        [0.0013, None,  None,  None,  31.18],
    ]
    add_sheet(wb, "granulo_sedimento_large", sed_headers, sed_rows)

    out_path = "atlas_import_example.xlsx"
    wb.save(out_path)
    print(f"✅ Fichier créé : {out_path}")

if __name__ == "__main__":
    main()
