"""
Script de canonisation des fichiers Excel bruts vers le format atlas_import_template.xlsx
Transforme les données géotechniques de différentes sources vers le format standardisé Atlas V3

Usage:
    python canonize_xlsx_to_atlas_import.py <input_xlsx> <output_xlsx> [--source "Nom Source"]
    python canonize_xlsx_to_atlas_import.py --batch <input_dir> <output_dir>

Auteur: Atlas Géotechnique
Version: 1.0.0
"""

import openpyxl
from openpyxl import Workbook
from pathlib import Path
import re
import argparse
from datetime import datetime
from typing import Dict, List, Tuple, Optional, Any
import json

CODE_MAPPING = {
    'BOHOU': 'BOHOU',
    'LAMA FEING': 'LAMA_FEING',
    'LAMA FIENG': 'LAMA_FEING',
    'LAMA TCHAMDE': 'LAMA_TCHAMDE',
    'LAMA TCHAMDÈ': 'LAMA_TCHAMDE',
    'KONSOGOU T1': 'KONSOGOUT1',
    'KONSOGOUT1': 'KONSOGOUT1',
    'KONSOGOU T2': 'KONSOGOUT2',
    'KONSOGOUT2': 'KONSOGOUT2',
    'NASSABLE': 'NASSABLE',
    'NASSABLÉ': 'NASSABLE',
    'NASSABL': 'NASSABLE',
    'KONTONGBONGUE': 'KONTONGBONGUE',
}


def extract_code_site(sheet_name: str) -> Optional[str]:
    if not sheet_name:
        return None

    name = str(sheet_name).strip().upper()
    name = re.sub(r'[_\-]+', ' ', name)
    name = re.sub(r'\s+', ' ', name)

    for k, v in CODE_MAPPING.items():
        if k in name:
            return v

    m = re.search(r'\(([^\)]+)\)', name)
    if m:
        inside = m.group(1)
        for k, v in CODE_MAPPING.items():
            if k in inside:
                return v

    return None

# ============================================================================
# CONFIGURATION
# ============================================================================

# Patterns pour détecter les types de feuilles
SHEET_PATTERNS = {
    'agt': r'AGT|tamisage|granulo.*tamisage',
    'ags': r'AGS|sedimentation|sedimento',
    'atterberg': r'limite.*att|atterberg|wl|wp',
    'vbs': r'bleu|vbs|methylene',
    'densite': r'densit[eé]|rho|masse.*vol',
    'teneur_eau': r'teneur.*eau|water.*content|humidite',
    'classification': r'classif|hrb|uscs|unified',
    'proctor': r'proctor|compactage',
}

# Mapping des profondeurs depuis les noms de colonnes/feuilles
DEPTH_PATTERNS = [
    (r'(\d+)[,.]?(\d*)\s*m', lambda m: float(f"{m.group(1)}.{m.group(2) or '0'}")),
    (r'(\d+)[,.]?(\d*)\s*\(?m\)?', lambda m: float(f"{m.group(1)}.{m.group(2) or '0'}")),
    (r'prof.*?(\d+)[,.]?(\d*)', lambda m: float(f"{m.group(1)}.{m.group(2) or '0'}")),
    (r'@(\d+)[,.]?(\d*)', lambda m: float(f"{m.group(1)}.{m.group(2) or '0'}")),
]

# ============================================================================
# CLASSES UTILITAIRES
# ============================================================================

class LocalityExtractor:
    """Extrait les localités depuis les noms de feuilles"""
    
    # Patterns connus pour les localités
    KNOWN_LOCALITIES = [
        'bohou', 'lama', 'feing', 'tchamde', 'katore', 'tsevie', 'deve',
        'assahoun', 'badja', 'keve', 'yade', 'tchitchao', 'konsogou',
        'nagerou', 'dapaong', 'nassable', 'kontongbongue'
    ]
    
    @classmethod
    def extract(cls, sheet_name: str) -> Optional[str]:
        """Extrait la localité depuis le nom de feuille"""
        name_lower = sheet_name.lower()
        
        # Nettoyer le nom
        name_clean = re.sub(r'agt|ags|all|\d+[,.]?\d*\s*m|_x\d+_', '', name_lower, flags=re.IGNORECASE)
        name_clean = re.sub(r'[_\-\(\)\[\]]', ' ', name_clean)
        name_clean = ' '.join(name_clean.split())
        
        # Chercher une localité connue
        for loc in cls.KNOWN_LOCALITIES:
            if loc in name_lower:
                return loc.capitalize()
        
        # Sinon, prendre le premier mot significatif
        words = [w for w in name_clean.split() if len(w) > 2]
        if words:
            return words[0].capitalize()
        
        return None


class DepthExtractor:
    """Extrait les profondeurs depuis les noms de colonnes/feuilles"""
    
    @classmethod
    def from_text(cls, text: str) -> Optional[float]:
        """Extrait une profondeur depuis un texte"""
        if not text:
            return None
        
        text_lower = str(text).lower()
        
        for pattern, converter in DEPTH_PATTERNS:
            match = re.search(pattern, text_lower)
            if match:
                try:
                    return converter(match)
                except:
                    pass
        
        return None
    
    @classmethod
    def from_column_header(cls, header: str) -> Optional[float]:
        """Extrait une profondeur depuis un en-tête de colonne"""
        return cls.from_text(header)


# ============================================================================
# PARSERS SPÉCIALISÉS
# ============================================================================

class GranuloParser:
    """Parse les données granulométriques (AGT/AGS)"""

    @staticmethod
    def _to_float(value: Any) -> Optional[float]:
        if value is None:
            return None
        if isinstance(value, (int, float)):
            try:
                return float(value)
            except Exception:
                return None

        try:
            s = str(value).strip()
            if not s:
                return None
            s = s.replace(' ', '').replace(',', '.')
            return float(s)
        except Exception:
            return None
    
    @classmethod
    def parse_sheet(cls, ws, sheet_name: str, source: str) -> List[Dict]:
        """Parse une feuille de granulométrie"""
        results = []
        
        # Détecter le type (tamisage ou sédimentation)
        is_ags = bool(re.search(r'ags|sediment', sheet_name, re.IGNORECASE))
        method = 'sedimentation' if is_ags else 'tamisage'
        
        # Extraire la localité
        locality = LocalityExtractor.extract(sheet_name)
        code_site = extract_code_site(sheet_name)
        
        # Lire les données
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 3:
            return results
        
        # Trouver la ligne d'en-tête (contient "tamis" ou "mm")
        header_row_idx = None
        for i, row in enumerate(rows):
            row_str = ' '.join(str(c or '').lower() for c in row)
            if 'tamis' in row_str or 'mm' in row_str or 'ouverture' in row_str:
                header_row_idx = i
                break
        
        if header_row_idx is None:
            return results
        
        header = rows[header_row_idx]
        
        # Identifier les colonnes de profondeur
        depth_cols = {}  # col_idx -> depth_m
        sieve_col = None
        
        for col_idx, cell in enumerate(header):
            cell_str = str(cell or '').lower()
            
            # Colonne tamis
            if 'tamis' in cell_str or 'mm' in cell_str or 'ouverture' in cell_str:
                sieve_col = col_idx
                continue
            
            # Colonnes de profondeur
            depth = DepthExtractor.from_column_header(str(cell or ''))
            if depth is not None:
                depth_cols[col_idx] = depth
            elif 'passant' in cell_str:
                # Chercher la profondeur dans le texte
                depth = DepthExtractor.from_text(cell_str)
                if depth:
                    depth_cols[col_idx] = depth
        
        # Cas 1: feuille "wide" (plusieurs profondeurs dans la même feuille)
        if sieve_col is not None and depth_cols:
            for row in rows[header_row_idx + 1:]:
                try:
                    sieve_val = row[sieve_col]
                    if sieve_val is None:
                        continue

                    sieve_mm = cls._to_float(sieve_val)
                    if sieve_mm is None:
                        continue
                    if sieve_mm <= 0 or sieve_mm > 100:
                        continue

                    for col_idx, depth_m in depth_cols.items():
                        passing_val = row[col_idx] if col_idx < len(row) else None
                        if passing_val is None:
                            continue

                        try:
                            passing_pct = cls._to_float(passing_val)
                            if passing_pct is None:
                                continue
                            if 0 <= passing_pct <= 100:
                                results.append({
                                    'code_site': code_site,
                                    'locality': locality,
                                    'depth_m': depth_m,
                                    'sieve_mm': sieve_mm,
                                    'passing_pct': passing_pct,
                                    'method': method,
                                    'source': source
                                })
                        except (ValueError, TypeError):
                            pass
                except (ValueError, TypeError, IndexError):
                    continue

            return results

        # Cas 2: feuille "long" par profondeur (une seule profondeur dans le nom de feuille)
        if sieve_col is None:
            return results

        depth_m = DepthExtractor.from_text(sheet_name)
        if depth_m is None:
            return results

        passing_col = None
        for col_idx, cell in enumerate(header):
            cell_str = str(cell or '').lower()
            if col_idx == sieve_col:
                continue
            if 'passant' in cell_str or '%' in cell_str:
                passing_col = col_idx
                break

        if passing_col is None:
            for col_idx, cell in enumerate(header):
                if col_idx == sieve_col:
                    continue
                if cell is not None:
                    passing_col = col_idx
                    break

        if passing_col is None:
            return results
        
        for row in rows[header_row_idx + 1:]:
            try:
                sieve_val = row[sieve_col]
                passing_val = row[passing_col] if passing_col < len(row) else None
                if sieve_val is None or passing_val is None:
                    continue

                sieve_mm = cls._to_float(sieve_val)
                if sieve_mm is None:
                    continue
                if sieve_mm <= 0 or sieve_mm > 100:
                    continue

                passing_pct = cls._to_float(passing_val)
                if passing_pct is None:
                    continue
                if 0 <= passing_pct <= 100:
                    results.append({
                        'code_site': code_site,
                        'locality': locality,
                        'depth_m': depth_m,
                        'sieve_mm': sieve_mm,
                        'passing_pct': passing_pct,
                        'method': method,
                        'source': source
                    })
            except (ValueError, TypeError, IndexError):
                continue
        
        return results


class AtterbergParser:
    """Parse les données Atterberg (WL, WP, IP)"""
    
    @classmethod
    def parse_sheet(cls, ws, sheet_name: str, source: str) -> List[Dict]:
        """Parse une feuille Atterberg"""
        results = []

        code_site = extract_code_site(sheet_name)
        
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            return results
        
        # Chercher l'en-tête
        header_row_idx = None
        for i, row in enumerate(rows):
            row_str = ' '.join(str(c or '').lower() for c in row)
            if 'wl' in row_str or 'wp' in row_str or 'liquidit' in row_str or 'plasticit' in row_str:
                header_row_idx = i
                break
        
        if header_row_idx is None:
            return results
        
        header = rows[header_row_idx]
        
        # Identifier les colonnes
        col_map = {}
        for col_idx, cell in enumerate(header):
            cell_str = str(cell or '').lower()
            if 'localit' in cell_str or 'site' in cell_str or 'code' in cell_str:
                col_map['locality'] = col_idx
            elif 'prof' in cell_str or 'depth' in cell_str:
                col_map['depth'] = col_idx
            elif 'wl' in cell_str or 'liquidit' in cell_str:
                col_map['wl'] = col_idx
            elif 'wp' in cell_str or 'plasticit' in cell_str:
                col_map['wp'] = col_idx
            elif 'ip' in cell_str or 'indice' in cell_str:
                col_map['ip'] = col_idx
        
        # Parser les données
        for row in rows[header_row_idx + 1:]:
            try:
                locality = row[col_map.get('locality', 0)] if 'locality' in col_map else LocalityExtractor.extract(sheet_name)
                depth_m = float(row[col_map['depth']]) if 'depth' in col_map and row[col_map['depth']] else None
                wl = float(row[col_map['wl']]) if 'wl' in col_map and row[col_map['wl']] else None
                wp = float(row[col_map['wp']]) if 'wp' in col_map and row[col_map['wp']] else None
                
                if depth_m is None:
                    continue

                if wl is not None and wp is not None:
                    results.append({
                        'code_site': code_site,
                        'locality': locality,
                        'depth_m': depth_m,
                        'wl': wl,
                        'wp': wp,
                        'source': source
                    })
            except (ValueError, TypeError, IndexError, KeyError):
                continue
        
        return results


class VBSParser:
    """Parse les données VBS (Bleu de méthylène)"""
    
    @classmethod
    def parse_sheet(cls, ws, sheet_name: str, source: str) -> List[Dict]:
        """Parse une feuille VBS"""
        results = []

        code_site = extract_code_site(sheet_name)
        
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            return results
        
        # Chercher l'en-tête
        header_row_idx = None
        for i, row in enumerate(rows):
            row_str = ' '.join(str(c or '').lower() for c in row)
            if 'vbs' in row_str or 'bleu' in row_str or 'methylene' in row_str:
                header_row_idx = i
                break
        
        if header_row_idx is None:
            return results
        
        header = rows[header_row_idx]
        
        # Identifier les colonnes
        col_map = {}
        for col_idx, cell in enumerate(header):
            cell_str = str(cell or '').lower()
            if 'localit' in cell_str or 'site' in cell_str or 'code' in cell_str:
                col_map['locality'] = col_idx
            elif 'prof' in cell_str or 'depth' in cell_str:
                col_map['depth'] = col_idx
            elif 'vbs' in cell_str or 'bleu' in cell_str or 'valeur' in cell_str:
                col_map['vbs'] = col_idx
        
        # Parser les données
        for row in rows[header_row_idx + 1:]:
            try:
                locality = row[col_map.get('locality', 0)] if 'locality' in col_map else LocalityExtractor.extract(sheet_name)
                depth_m = float(row[col_map['depth']]) if 'depth' in col_map and row[col_map['depth']] else None
                vbs = float(row[col_map['vbs']]) if 'vbs' in col_map and row[col_map['vbs']] else None
                
                if depth_m is None:
                    continue

                if vbs is not None:
                    results.append({
                        'code_site': code_site,
                        'locality': locality,
                        'depth_m': depth_m,
                        'vbs': vbs,
                        'source': source
                    })
            except (ValueError, TypeError, IndexError, KeyError):
                continue
        
        return results


# ============================================================================
# GÉNÉRATEUR DE FICHIER CANONIQUE
# ============================================================================

class AtlasImportGenerator:
    """Génère un fichier atlas_import.xlsx à partir des données parsées"""
    
    def __init__(self, source: str):
        self.source = source
        self.sondages: Dict[str, Dict] = {}  # code -> {locality, depths, ...}
        self.echantillons: List[Dict] = []
        self.atterberg: List[Dict] = []
        self.vbs: List[Dict] = []
        self.granulo: List[Dict] = []
    
    def add_granulo_data(self, data: List[Dict]):
        """Ajoute des données granulométriques"""
        for item in data:
            locality = item.get('locality', 'Unknown')
            depth_m = item.get('depth_m')

            # Créer/mettre à jour le sondage
            code = item.get('code_site') or self._get_sondage_code(locality)
            if code not in self.sondages:
                self.sondages[code] = {
                    'locality': locality,
                    'depths': set()
                }
            if depth_m:
                self.sondages[code]['depths'].add(depth_m)
            
            # Ajouter le point granulo
            self.granulo.append({
                'code': code,
                'depth_m': depth_m,
                'sieve_mm': item.get('sieve_mm'),
                'passing_pct': item.get('passing_pct'),
                'method': item.get('method')
            })
    
    def add_atterberg_data(self, data: List[Dict]):
        """Ajoute des données Atterberg"""
        for item in data:
            locality = item.get('locality', 'Unknown')
            depth_m = item.get('depth_m')

            code = item.get('code_site') or self._get_sondage_code(locality)
            if code not in self.sondages:
                self.sondages[code] = {
                    'locality': locality,
                    'depths': set()
                }
            if depth_m:
                self.sondages[code]['depths'].add(depth_m)
            
            self.atterberg.append({
                'code': code,
                'depth_m': depth_m,
                'wl': item.get('wl'),
                'wp': item.get('wp')
            })
    
    def add_vbs_data(self, data: List[Dict]):
        """Ajoute des données VBS"""
        for item in data:
            locality = item.get('locality', 'Unknown')
            depth_m = item.get('depth_m')

            code = item.get('code_site') or self._get_sondage_code(locality)
            if code not in self.sondages:
                self.sondages[code] = {
                    'locality': locality,
                    'depths': set()
                }
            if depth_m:
                self.sondages[code]['depths'].add(depth_m)
            
            self.vbs.append({
                'code': code,
                'depth_m': depth_m,
                'vbs': item.get('vbs')
            })
    
    def _get_sondage_code(self, locality: str) -> str:
        """Génère un code sondage depuis la localité"""
        if not locality:
            return 'UNKNOWN-S1'
        
        # Normaliser
        code = locality.upper().replace(' ', '-')
        code = re.sub(r'[^A-Z0-9\-]', '', code)
        
        return f"{code}-S1"
    
    def generate_xlsx(self, output_path: Path):
        """Génère le fichier Excel canonique"""
        wb = Workbook()
        wb.remove(wb.active)
        
        today = datetime.now().strftime('%Y-%m-%d')
        
        # Feuille sondages
        ws_sondages = wb.create_sheet('sondages')
        ws_sondages.append(['code_site', 'localite', 'date', 'adm3', 'adm2', 'lat', 'lon', 'source'])
        for code, info in self.sondages.items():
            ws_sondages.append([code, info['locality'], today, None, None, None, None, self.source])
        
        # Feuille echantillons
        ws_echantillons = wb.create_sheet('echantillons')
        ws_echantillons.append(['code_site', 'depth_m', 'date', 'laboratory', 'rho_s_gcm3', 'water_content_w', 'is_index', 'commentaire'])
        
        # Générer les échantillons depuis les profondeurs collectées
        for code, info in self.sondages.items():
            for depth in sorted(info['depths']):
                ws_echantillons.append([code, depth, today, self.source, None, None, None, None])
        
        # Feuille atterberg
        ws_atterberg = wb.create_sheet('atterberg')
        ws_atterberg.append(['code_site', 'depth_m', 'wl', 'wp'])
        for item in self.atterberg:
            ws_atterberg.append([item['code'], item['depth_m'], item['wl'], item['wp']])
        
        # Feuille vbs
        ws_vbs = wb.create_sheet('vbs')
        ws_vbs.append(['code_site', 'depth_m', 'vbs', 'commentaire'])
        for item in self.vbs:
            ws_vbs.append([item['code'], item['depth_m'], item['vbs'], None])
        
        # Feuille proctor (vide)
        ws_proctor = wb.create_sheet('proctor')
        ws_proctor.append(['code_site', 'depth_m', 'proctor_type', 'gamma_d_max', 'w_opt'])
        
        # Feuille granulo (format long)
        ws_granulo = wb.create_sheet('granulo')
        ws_granulo.append(['code_site', 'depth_m', 'sieve_mm', 'passing_pct', 'method'])
        for item in self.granulo:
            ws_granulo.append([item['code'], item['depth_m'], item['sieve_mm'], item['passing_pct'], item['method']])
        
        # Sauvegarder
        wb.save(output_path)
        print(f"[OK] Fichier généré: {output_path}")
        print(f"     - {len(self.sondages)} sondages")
        print(f"     - {sum(len(s['depths']) for s in self.sondages.values())} échantillons")
        print(f"     - {len(self.atterberg)} essais Atterberg")
        print(f"     - {len(self.vbs)} essais VBS")
        print(f"     - {len(self.granulo)} points granulo")


# ============================================================================
# FONCTION PRINCIPALE
# ============================================================================

def canonize_xlsx(input_path: Path, output_path: Path, source: str = None):
    """Canonise un fichier Excel brut vers le format atlas_import"""
    
    if source is None:
        source = input_path.stem
    
    print(f"\n[*] Canonisation de: {input_path}")
    print(f"    Source: {source}")
    
    try:
        wb = openpyxl.load_workbook(input_path, data_only=True)
    except Exception as e:
        print(f"[ERREUR] Impossible d'ouvrir le fichier: {e}")
        return False
    
    generator = AtlasImportGenerator(source)
    
    # Parser chaque feuille
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_lower = sheet_name.lower()
        
        print(f"    Feuille: {sheet_name}")
        
        # Détecter le type de feuille
        if re.search(SHEET_PATTERNS['agt'], sheet_name, re.IGNORECASE) or re.search(
            SHEET_PATTERNS['ags'], sheet_name, re.IGNORECASE
        ):
            data = GranuloParser.parse_sheet(ws, sheet_name, source)
            if data:
                generator.add_granulo_data(data)
                print(f"      -> {len(data)} points granulo")
        
        elif re.search(SHEET_PATTERNS['atterberg'], sheet_name, re.IGNORECASE):
            data = AtterbergParser.parse_sheet(ws, sheet_name, source)
            if data:
                generator.add_atterberg_data(data)
                print(f"      -> {len(data)} essais Atterberg")
        
        elif re.search(SHEET_PATTERNS['vbs'], sheet_name, re.IGNORECASE):
            data = VBSParser.parse_sheet(ws, sheet_name, source)
            if data:
                generator.add_vbs_data(data)
                print(f"      -> {len(data)} essais VBS")
    
    # Générer le fichier de sortie
    generator.generate_xlsx(output_path)
    
    return True


def batch_canonize(input_dir: Path, output_dir: Path):
    """Canonise tous les fichiers Excel d'un répertoire"""
    
    output_dir.mkdir(parents=True, exist_ok=True)
    
    xlsx_files = list(input_dir.glob('*.xlsx'))
    print(f"\n[*] Batch canonisation: {len(xlsx_files)} fichiers trouvés")
    
    success = 0
    for xlsx_file in xlsx_files:
        if xlsx_file.name.startswith('~$'):  # Fichiers temporaires
            continue
        
        output_file = output_dir / f"atlas_import_{xlsx_file.stem}.xlsx"
        if canonize_xlsx(xlsx_file, output_file):
            success += 1
    
    print(f"\n[RÉSUMÉ] {success}/{len(xlsx_files)} fichiers canonisés avec succès")


# ============================================================================
# CLI
# ============================================================================

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Canonise les fichiers Excel vers le format atlas_import')
    parser.add_argument('input', help='Fichier ou répertoire d\'entrée')
    parser.add_argument('output', help='Fichier ou répertoire de sortie')
    parser.add_argument('--source', help='Nom de la source (défaut: nom du fichier)')
    parser.add_argument('--batch', action='store_true', help='Mode batch (répertoires)')
    
    args = parser.parse_args()
    
    input_path = Path(args.input)
    output_path = Path(args.output)
    
    if args.batch:
        batch_canonize(input_path, output_path)
    else:
        canonize_xlsx(input_path, output_path, args.source)
