"""
Import spécifique pour le fichier NICABOU Ninsao Vianney
Format différent des autres sources - feuilles LIMITE ATT, BLEU, AGT/AGS par localité

Usage:
    python import_nicabou.py [--dry-run]
"""

import openpyxl
import psycopg2
from pathlib import Path
import argparse
import re
from typing import Dict, List, Optional
import uuid

DB_CONFIG = {
    'host': 'localhost',
    'port': 5432,
    'database': 'atlas_clean',
    'user': 'atlas',
    'password': 'atlas'
}

XLSX_PATH = Path(__file__).parent.parent / 'data' / 'xlsx' / 'NICABOU Ninsao Vianney.xlsx'
SOURCE = 'NICABOU Ninsao Vianney'

# Mapping localités Excel -> codes en base
LOCALITE_MAPPING = {
    'Apéhémé': 'APEHEME',
    'Apeheme': 'APEHEME',
    'APEHEME': 'APEHEME',
    'Dzogbécopé': 'DZOGBECOPE',
    'Dzogbecope': 'DZOGBECOPE',
    'DZOGBECOPE': 'DZOGBECOPE',
    'Davié': 'DAVIE',
    'Davie': 'DAVIE',
    'DAVIE': 'DAVIE',
    'Tekpo': 'TEKPO',  # Pas encore créé
    'TEKPO': 'TEKPO',
}

def parse_depth(depth_str: str) -> Optional[float]:
    """Parse une profondeur comme '1m', '1,5m', '2m' en float"""
    if not depth_str:
        return None
    depth_str = str(depth_str).strip().lower()
    if depth_str == 'moyenne':
        return None
    # Enlever le 'm' et convertir
    depth_str = depth_str.replace('m', '').replace(',', '.').strip()
    try:
        return float(depth_str)
    except:
        return None

def get_sondage_id(conn, code: str) -> Optional[str]:
    """Récupère l'ID du sondage par son code"""
    cur = conn.cursor()
    cur.execute("SELECT id FROM atlas.sondages WHERE code = %s", (code,))
    row = cur.fetchone()
    return str(row[0]) if row else None

def get_or_create_echantillon(conn, sondage_id: str, depth_m: float, source: str) -> str:
    """Récupère ou crée un échantillon"""
    cur = conn.cursor()
    cur.execute("""
        SELECT id FROM atlas.echantillons 
        WHERE sondage_id = %s AND depth_m = %s
    """, (sondage_id, depth_m))
    row = cur.fetchone()
    if row:
        return str(row[0])
    
    # Créer l'échantillon
    ech_id = str(uuid.uuid4())
    cur.execute("""
        INSERT INTO atlas.echantillons (id, sondage_id, depth_m, laboratory)
        VALUES (%s, %s, %s, %s)
    """, (ech_id, sondage_id, depth_m, source))
    return ech_id

def import_atterberg(conn, wb, dry_run: bool) -> int:
    """Importe les données Atterberg depuis la feuille LIMITE ATT"""
    if 'LIMITE ATT' not in wb.sheetnames:
        print("  ⚠️  Feuille LIMITE ATT non trouvée")
        return 0
    
    ws = wb['LIMITE ATT']
    count = 0
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        # Format: _, LOCALITE, PROFONDEUR, WL, WP, IP, IC
        if len(row) < 6:
            continue
        localite = str(row[1]).strip() if row[1] else None
        depth_str = str(row[2]).strip() if row[2] else None
        wl = row[3]
        wp = row[4]
        ip = row[5]
        
        if not localite or not depth_str:
            continue
        
        depth_m = parse_depth(depth_str)
        if depth_m is None:
            continue
        
        # Mapper la localité vers le code
        code = LOCALITE_MAPPING.get(localite)
        if not code:
            continue
        
        sondage_id = get_sondage_id(conn, code)
        if not sondage_id:
            continue
        
        ech_id = get_or_create_echantillon(conn, sondage_id, depth_m, SOURCE)
        
        # Insérer l'essai Atterberg
        cur = conn.cursor()
        cur.execute("""
            INSERT INTO atlas.essais_atterberg (id, echantillon_id, wl, wp, ip_generated)
            VALUES (%s, %s, %s, %s, %s)
            ON CONFLICT DO NOTHING
        """, (str(uuid.uuid4()), ech_id, wl, wp, ip))
        count += 1
    
    return count

def import_vbs(conn, wb, dry_run: bool) -> int:
    """Importe les données VBS depuis la feuille BLEU"""
    if 'BLEU' not in wb.sheetnames:
        print("  ⚠️  Feuille BLEU non trouvée")
        return 0
    
    ws = wb['BLEU']
    count = 0
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        # Format: Référence, Localité, profondeurs, VBS, Interprétation
        if len(row) < 4:
            continue
        localite = str(row[1]).strip() if row[1] else None
        depth_m = row[2]
        vbs = row[3]
        
        if not localite or depth_m is None or vbs is None:
            continue
        
        try:
            depth_m = float(depth_m)
            vbs = float(vbs)
        except:
            continue
        
        code = LOCALITE_MAPPING.get(localite)
        if not code:
            continue
        
        sondage_id = get_sondage_id(conn, code)
        if not sondage_id:
            continue
        
        ech_id = get_or_create_echantillon(conn, sondage_id, depth_m, SOURCE)
        
        cur = conn.cursor()
        cur.execute("""
            INSERT INTO atlas.essais_vbs (id, echantillon_id, vbs)
            VALUES (%s, %s, %s)
            ON CONFLICT DO NOTHING
        """, (str(uuid.uuid4()), ech_id, vbs))
        count += 1
    
    return count

def import_granulo(conn, wb, dry_run: bool) -> int:
    """Importe les données granulométriques depuis les feuilles AGT/AGS"""
    count = 0
    
    # Feuilles AGT (tamisage) et AGS (sédimentation)
    for sheet_name in wb.sheetnames:
        if not (sheet_name.startswith('AGT ') or sheet_name.startswith('AGS ')):
            continue
        
        ws = wb[sheet_name]
        is_sedimentation = sheet_name.startswith('AGS ')
        method = 'sedimentation' if is_sedimentation else 'tamisage'
        
        # Extraire la localité et profondeur du nom de feuille
        # Format: "AGT Apeheme" ou "AGT Apéhémé 1m all"
        parts = sheet_name.split(' ', 1)
        if len(parts) < 2:
            continue
        rest = parts[1].strip()
        
        # Vérifier si c'est le format "Localité Xm all"
        depth_from_name = None
        localite_raw = rest
        match = re.match(r'^(.+?)\s+(\d+(?:,\d+)?)\s*m\s*all$', rest, re.IGNORECASE)
        if match:
            localite_raw = match.group(1).strip()
            depth_str = match.group(2).replace(',', '.')
            depth_from_name = float(depth_str)
        
        # Mapper la localité
        code = None
        for key, val in LOCALITE_MAPPING.items():
            if key.lower() == localite_raw.lower():
                code = val
                break
        
        if not code:
            continue
        
        sondage_id = get_sondage_id(conn, code)
        if not sondage_id:
            continue
        
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            continue
        
        # Si profondeur dans le nom, format simple: Tamis | Passant
        if depth_from_name is not None:
            # Chercher la colonne Tamis et Passant
            tamis_col = None
            passant_col = None
            
            for row in rows:
                for col_idx, cell in enumerate(row):
                    if cell:
                        cell_str = str(cell).lower()
                        if 'tamis' in cell_str:
                            tamis_col = col_idx
                        if 'pas' in cell_str and 'cum' in cell_str:
                            passant_col = col_idx
            
            if tamis_col is None:
                tamis_col = 3  # Default basé sur le format observé
            if passant_col is None:
                passant_col = 7  # Default "Pas. Cum. (%)"
            
            for row in rows:
                if len(row) <= max(tamis_col, passant_col):
                    continue
                
                tamis_val = row[tamis_col]
                passant_val = row[passant_col]
                
                if tamis_val is None or passant_val is None:
                    continue
                
                try:
                    sieve_mm = float(tamis_val)
                    passant_pct = float(passant_val)
                except:
                    continue
                
                ech_id = get_or_create_echantillon(conn, sondage_id, depth_from_name, SOURCE)
                
                cur = conn.cursor()
                cur.execute("""
                    INSERT INTO atlas.granulo_points (id, echantillon_id, sieve_mm, passing_pct, method)
                    VALUES (%s, %s, %s, %s, %s)
                    ON CONFLICT DO NOTHING
                """, (str(uuid.uuid4()), ech_id, sieve_mm, passant_pct, method))
                count += 1
        else:
            # Format wide avec colonnes par profondeur
            header_row = rows[0]
            depth_cols = {}  # col_idx -> depth_m
            
            for col_idx, cell in enumerate(header_row):
                if cell:
                    cell_str = str(cell).lower()
                    if '1m' in cell_str and '1,5' not in cell_str and '1.5' not in cell_str:
                        depth_cols[col_idx] = 1.0
                    elif '1,5m' in cell_str or '1.5m' in cell_str:
                        depth_cols[col_idx] = 1.5
                    elif '2m' in cell_str:
                        depth_cols[col_idx] = 2.0
            
            if not depth_cols:
                continue
            
            tamis_col = 0
            for col_idx, cell in enumerate(header_row):
                if cell and 'tamis' in str(cell).lower():
                    tamis_col = col_idx
                    break
            
            for row in rows[1:]:
                if len(row) <= tamis_col:
                    continue
                
                tamis_val = row[tamis_col]
                if tamis_val is None:
                    continue
                
                try:
                    sieve_mm = float(tamis_val)
                except:
                    continue
                
                for col_idx, depth_m in depth_cols.items():
                    if col_idx >= len(row):
                        continue
                    passant = row[col_idx]
                    if passant is None:
                        continue
                    try:
                        passant_pct = float(passant)
                    except:
                        continue
                    
                    ech_id = get_or_create_echantillon(conn, sondage_id, depth_m, SOURCE)
                    
                    cur = conn.cursor()
                    cur.execute("""
                        INSERT INTO atlas.granulo_points (id, echantillon_id, sieve_mm, passing_pct, method)
                        VALUES (%s, %s, %s, %s, %s)
                        ON CONFLICT DO NOTHING
                    """, (str(uuid.uuid4()), ech_id, sieve_mm, passant_pct, method))
                    count += 1
    
    return count

def main():
    parser = argparse.ArgumentParser(description='Import NICABOU data')
    parser.add_argument('--dry-run', action='store_true', help='Dry run mode')
    args = parser.parse_args()
    
    print(f"\n{'[DRY-RUN] ' if args.dry_run else ''}Import NICABOU Ninsao Vianney")
    print(f"  Fichier: {XLSX_PATH}")
    
    if not XLSX_PATH.exists():
        print(f"  ❌ Fichier non trouvé: {XLSX_PATH}")
        return
    
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    conn = psycopg2.connect(**DB_CONFIG)
    
    try:
        atterberg_count = import_atterberg(conn, wb, args.dry_run)
        vbs_count = import_vbs(conn, wb, args.dry_run)
        granulo_count = import_granulo(conn, wb, args.dry_run)
        
        print(f"  ✅ Atterberg: {atterberg_count}")
        print(f"  ✅ VBS: {vbs_count}")
        print(f"  ✅ Granulo: {granulo_count}")
        
        if args.dry_run:
            conn.rollback()
            print("  🔄 Rollback (dry-run)")
        else:
            conn.commit()
            print("  💾 Commit OK")
    except Exception as e:
        conn.rollback()
        print(f"  ❌ Erreur: {e}")
        raise
    finally:
        conn.close()
        wb.close()

if __name__ == '__main__':
    main()
