"""
Script de conversion batch optimisé pour fichiers Excel
Extrait les données et génère:
- CSV par feuille
- Fichier Markdown compilé avec toutes les feuilles
"""

import openpyxl
from pathlib import Path
import csv
from datetime import datetime
import re

# Configuration
XLSX_DIR = Path("data/xlsx")
OUTPUT_DIR = Path("data/xlsx_convert_md")
EXCLUDE_DIR = "IMPORT"

# Dictionnaire de traduction FR -> EN
TRANSLATIONS = {
    # Colonnes communes
    'code_site': 'site_code',
    'localite': 'locality',
    'date': 'date',
    'profondeur': 'depth_m',
    'depth_m': 'depth_m',
    'laboratoire': 'laboratory',
    'laboratory': 'laboratory',

    # Atterberg
    'wl': 'liquid_limit',
    'wp': 'plastic_limit',
    'ip': 'plasticity_index',
    'limite_liquidite': 'liquid_limit',
    'limite_plasticite': 'plastic_limit',
    'indice_plasticite': 'plasticity_index',

    # VBS
    'vbs': 'methylene_blue_value',
    'bleu_methylene': 'methylene_blue_value',
    'commentaire': 'comment',
    'caracteristique': 'characteristic',

    # Proctor
    'proctor_type': 'proctor_type',
    'gamma_d_max': 'max_dry_density',
    'w_opt': 'optimum_water_content',
    'densite_seche_max': 'max_dry_density',
    'teneur_eau_opt': 'optimum_water_content',

    # Densité
    'rho_s_gcm3': 'absolute_density',
    'rho_d_app': 'apparent_density',
    'densite_absolue': 'absolute_density',
    'densite_apparente': 'apparent_density',
    'water_content_w': 'water_content',
    'teneur_eau': 'water_content',

    # Granulo
    'sieve_mm': 'sieve_mm',
    'tamis_mm': 'sieve_mm',
    'ouverture_tamis': 'sieve_mm',
    'passant': 'passing_pct',
    'passants': 'passing_pct',
    'passant_pct': 'passing_pct',
    'refus': 'retained_pct',
    'refus_pct': 'retained_pct',
    'masse_refus': 'mass_retained_g',
    'mass_refus_cum_g': 'cumulative_mass_retained_g',
    'refus_cum_pct': 'cumulative_retained_pct',

    # Classification
    'hrb': 'hrb_classification',
    'unified': 'uscs_classification',
    'bm': 'methylene_blue_classification',
    'classification_hrb': 'hrb_classification',
    'classification_uscs': 'uscs_classification',

    # Divers
    'adm1': 'region',
    'adm2': 'prefecture',
    'adm3': 'commune',
    'lat': 'latitude',
    'lon': 'longitude',
    'source': 'source',
    'note': 'note',
    'is_index': 'swelling_index',
    'eg': 'swelling_eg',
    'wi': 'swelling_wi',
}

def normalize_header(header):
    """Normalise un en-tête de colonne"""
    if not header:
        return None

    # Convertir en string et nettoyer
    header = str(header).strip().lower()

    # Supprimer caractères spéciaux
    header = re.sub(r'[^\w\s]', '', header)

    # Remplacer espaces par underscore
    header = re.sub(r'\s+', '_', header)

    return header

def translate_header(header):
    """Traduit un en-tête FR -> EN"""
    normalized = normalize_header(header)
    if not normalized:
        return header

    # Chercher traduction exacte
    if normalized in TRANSLATIONS:
        return TRANSLATIONS[normalized]

    # Chercher traduction partielle
    for fr, en in TRANSLATIONS.items():
        if fr in normalized:
            return en

    # Retourner normalisé si pas de traduction
    return normalized

def export_sheet_to_csv(sheet, output_path):
    """Exporte une feuille Excel vers CSV et retourne les données"""
    rows = list(sheet.iter_rows(values_only=True))

    if not rows or len(rows) < 2:
        return None, 0

    # Première ligne = en-têtes
    headers = rows[0]
    translated_headers = [translate_header(h) if h is not None else f'col_{i}'
                         for i, h in enumerate(headers)]

    # Filtrer lignes non-vides
    data_rows = []
    for row in rows[1:]:
        if any(cell is not None and str(cell).strip() for cell in row):
            data_rows.append(row)

    if not data_rows:
        return None, 0

    # Écrire CSV
    with open(output_path, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(translated_headers)
        writer.writerows(data_rows)

    return {
        'headers': translated_headers,
        'rows': data_rows
    }, len(data_rows)

def format_table_markdown(headers, rows, max_rows=20):
    """Formate un tableau en Markdown"""
    md_lines = []

    # En-têtes
    md_lines.append("| " + " | ".join(str(h) for h in headers) + " |")
    md_lines.append("|" + "|".join(["---" for _ in headers]) + "|")

    # Données (limiter à max_rows pour la lisibilité)
    display_rows = rows[:max_rows]
    for row in display_rows:
        cells = [str(cell if cell is not None else '') for cell in row]
        md_lines.append("| " + " | ".join(cells) + " |")

    # Indication si tronqué
    if len(rows) > max_rows:
        md_lines.append(f"\n*({len(rows) - max_rows} lignes supplémentaires masquées)*\n")

    return "\n".join(md_lines)

def generate_markdown_report(file_stem, sheets_data, source_filename):
    """Génère un rapport Markdown compilé"""
    md_content = []

    # En-tête du document
    md_content.append(f"# {file_stem}")
    md_content.append(f"\n**Auteur/Source:** {source_filename}")
    md_content.append(f"\n**Date de conversion:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    md_content.append(f"\n**Nombre de feuilles:** {len(sheets_data)}")
    md_content.append("\n---\n")

    # Table des matières
    md_content.append("## Table des matières\n")
    for i, (sheet_name, _, _) in enumerate(sheets_data, 1):
        # Créer ancre pour lien interne
        anchor = re.sub(r'[^\w\s-]', '', sheet_name.lower())
        anchor = re.sub(r'\s+', '-', anchor)
        md_content.append(f"{i}. [{sheet_name}](#{anchor})")
    md_content.append("\n---\n")

    # Contenu de chaque feuille
    for sheet_name, data, row_count in sheets_data:
        # Ancre pour navigation
        anchor = re.sub(r'[^\w\s-]', '', sheet_name.lower())
        anchor = re.sub(r'\s+', '-', anchor)

        md_content.append(f"\n## {sheet_name} {{#{anchor}}}")
        md_content.append(f"\n**Lignes de données:** {row_count}\n")

        if data:
            table_md = format_table_markdown(data['headers'], data['rows'])
            md_content.append(table_md)

        md_content.append("\n---\n")

    return "\n".join(md_content)

def convert_xlsx_file(input_path, output_dir):
    """Convertit un fichier Excel en CSV + Markdown"""
    try:
        print(f"\n[*] Traitement: {input_path.name}")

        # Charger le workbook
        wb = openpyxl.load_workbook(input_path, data_only=True)

        # Créer sous-dossier pour ce fichier
        file_stem = input_path.stem
        file_output_dir = output_dir / file_stem
        file_output_dir.mkdir(parents=True, exist_ok=True)

        sheets_data = []
        total_rows = 0
        sheets_exported = 0

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]

            # Nettoyer nom de feuille pour nom de fichier
            safe_sheet_name = re.sub(r'[^\w\s-]', '', sheet_name)
            safe_sheet_name = re.sub(r'\s+', '_', safe_sheet_name)

            output_path = file_output_dir / f"{safe_sheet_name}.csv"

            data, row_count = export_sheet_to_csv(sheet, output_path)

            if row_count > 0:
                total_rows += row_count
                sheets_exported += 1
                sheets_data.append((sheet_name, data, row_count))
                print(f"   [OK] {sheet_name}: {row_count} lignes")
            else:
                # Supprimer fichier vide
                output_path.unlink(missing_ok=True)

        # Générer le Markdown compilé
        if sheets_data:
            md_path = file_output_dir / f"{file_stem}.md"
            md_content = generate_markdown_report(file_stem, sheets_data, input_path.name)

            with open(md_path, 'w', encoding='utf-8') as f:
                f.write(md_content)

            print(f"   [MD] Rapport Markdown créé: {md_path.name}")

        # Créer fichier manifest
        manifest_path = file_output_dir / "_manifest.txt"
        with open(manifest_path, 'w', encoding='utf-8') as f:
            f.write(f"Source: {input_path.name}\n")
            f.write(f"Auteur: {file_stem}\n")
            f.write(f"Conversion: {datetime.now().isoformat()}\n")
            f.write(f"Sheets: {sheets_exported}\n")
            f.write(f"Total rows: {total_rows}\n")
            f.write(f"\nFichiers générés:\n")
            f.write(f"- {sheets_exported} fichiers CSV\n")
            f.write(f"- 1 rapport Markdown compilé ({file_stem}.md)\n")

        print(f"   [RÉSUMÉ] {sheets_exported} feuille(s), {total_rows} lignes totales")

        return True

    except Exception as e:
        print(f"   [ERREUR] {input_path.name}: {e}")
        return False

def find_xlsx_files(xlsx_dir, exclude_dir):
    """Trouve tous les fichiers Excel sauf ceux dans exclude_dir"""
    xlsx_files = []

    for file_path in xlsx_dir.rglob("*.xlsx"):
        # Exclure les fichiers temporaires Excel
        if file_path.name.startswith('~$'):
            continue

        # Exclure le répertoire import
        if exclude_dir in str(file_path):
            continue

        xlsx_files.append(file_path)

    return sorted(xlsx_files)

def main():
    """Fonction principale"""
    print("="*70)
    print("CONVERSION FICHIERS EXCEL -> CSV + MARKDOWN")
    print("="*70)

    # Vérifier répertoire source
    if not XLSX_DIR.exists():
        print(f"\n[ERREUR] Répertoire source introuvable: {XLSX_DIR}")
        return

    # Créer répertoire de sortie
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    print(f"\n[*] Répertoire source: {XLSX_DIR}")
    print(f"[*] Répertoire sortie: {OUTPUT_DIR}")
    print(f"[*] Format: CSV + Markdown compilé")
    print(f"[*] Exclusion: {EXCLUDE_DIR}")

    # Trouver fichiers Excel
    xlsx_files = find_xlsx_files(XLSX_DIR, EXCLUDE_DIR)

    if not xlsx_files:
        print(f"\n[WARN] Aucun fichier Excel trouvé dans {XLSX_DIR}")
        return

    print(f"\n[*] {len(xlsx_files)} fichier(s) Excel trouvé(s)")

    # Traiter chaque fichier
    success_count = 0
    error_count = 0

    for xlsx_file in xlsx_files:
        if convert_xlsx_file(xlsx_file, OUTPUT_DIR):
            success_count += 1
        else:
            error_count += 1

    # Résumé
    print("\n" + "="*70)
    print("RÉSUMÉ")
    print("="*70)
    print(f"\nFichiers traités: {len(xlsx_files)}")
    print(f"  [OK] Succès: {success_count}")
    print(f"  [ERREUR] Échecs: {error_count}")
    print(f"\nRépertoire de sortie: {OUTPUT_DIR}")

    # Calculer statistiques
    csv_count = len(list(OUTPUT_DIR.rglob("*.csv")))
    md_count = len(list(OUTPUT_DIR.rglob("*.md")))

    total_size = 0
    for file in OUTPUT_DIR.rglob("*"):
        if file.is_file():
            total_size += file.stat().st_size

    print(f"\n[*] Fichiers créés:")
    print(f"    - {csv_count} fichiers CSV")
    print(f"    - {md_count} rapports Markdown")
    print(f"    - {success_count} manifests")
    print(f"\n[*] Taille totale: {total_size / 1024:.1f} KB")

    print("\n[OK] Conversion terminée!")
    print(f"\nStructure par dossier:")
    print(f"  - Fichiers CSV individuels par feuille")
    print(f"  - Fichier Markdown compilé avec toutes les feuilles")
    print(f"  - Fichier _manifest.txt avec métadonnées")

if __name__ == '__main__':
    main()
