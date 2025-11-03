"""
Script de conversion batch des fichiers Excel
Extrait et compile les données de toutes les feuilles Excel
Traduit les en-têtes et sauvegarde dans /data/xlsx_convert
"""

import openpyxl
from pathlib import Path
import json
from datetime import datetime
import re

# Configuration
XLSX_DIR = Path("data/xlsx")
OUTPUT_DIR = Path("data/xlsx_convert")
EXCLUDE_DIR = "IMPORT"  # Correspondre à la casse exacte du dossier

# Dictionnaire de traduction FR -> EN
TRANSLATIONS = {
    # Colonnes communes
    'code_site': 'site_code',
    'localite': 'locality',
    'date': 'date',
    'profondeur': 'depth_m',
    'depth_m': 'depth_m',
    'laboratoire': 'laboratory',
    'laboratory': 'laboratory',

    # Atterberg
    'wl': 'liquid_limit',
    'wp': 'plastic_limit',
    'ip': 'plasticity_index',
    'limite_liquidite': 'liquid_limit',
    'limite_plasticite': 'plastic_limit',
    'indice_plasticite': 'plasticity_index',

    # VBS
    'vbs': 'methylene_blue_value',
    'bleu_methylene': 'methylene_blue_value',
    'commentaire': 'comment',
    'caracteristique': 'characteristic',

    # Proctor
    'proctor_type': 'proctor_type',
    'gamma_d_max': 'max_dry_density',
    'w_opt': 'optimum_water_content',
    'densite_seche_max': 'max_dry_density',
    'teneur_eau_opt': 'optimum_water_content',

    # Densité
    'rho_s_gcm3': 'absolute_density',
    'rho_d_app': 'apparent_density',
    'densite_absolue': 'absolute_density',
    'densite_apparente': 'apparent_density',
    'water_content_w': 'water_content',
    'teneur_eau': 'water_content',

    # Granulo
    'sieve_mm': 'sieve_mm',
    'tamis_mm': 'sieve_mm',
    'ouverture_tamis': 'sieve_mm',
    'passant': 'passing_pct',
    'passants': 'passing_pct',
    'passant_pct': 'passing_pct',
    'refus': 'retained_pct',
    'refus_pct': 'retained_pct',
    'masse_refus': 'mass_retained_g',
    'mass_refus_cum_g': 'cumulative_mass_retained_g',
    'refus_cum_pct': 'cumulative_retained_pct',

    # Classification
    'hrb': 'hrb_classification',
    'unified': 'uscs_classification',
    'bm': 'methylene_blue_classification',
    'classification_hrb': 'hrb_classification',
    'classification_uscs': 'uscs_classification',

    # Divers
    'adm1': 'region',
    'adm2': 'prefecture',
    'adm3': 'commune',
    'lat': 'latitude',
    'lon': 'longitude',
    'source': 'source',
    'note': 'note',
    'is_index': 'swelling_index',
    'eg': 'swelling_eg',
    'wi': 'swelling_wi',
}

def normalize_header(header):
    """Normalise un en-tête de colonne"""
    if not header:
        return None

    # Convertir en string et nettoyer
    header = str(header).strip().lower()

    # Supprimer caractères spéciaux
    header = re.sub(r'[^\w\s]', '', header)

    # Remplacer espaces par underscore
    header = re.sub(r'\s+', '_', header)

    return header

def translate_header(header):
    """Traduit un en-tête FR -> EN"""
    normalized = normalize_header(header)
    if not normalized:
        return header

    # Chercher traduction exacte
    if normalized in TRANSLATIONS:
        return TRANSLATIONS[normalized]

    # Chercher traduction partielle
    for fr, en in TRANSLATIONS.items():
        if fr in normalized:
            return en

    # Retourner normalisé si pas de traduction
    return normalized

def extract_sheet_data(sheet):
    """Extrait toutes les données d'une feuille Excel"""
    data = {
        'name': sheet.title,
        'headers': [],
        'rows': [],
        'translated_headers': []
    }

    rows = list(sheet.iter_rows(values_only=True))

    if not rows:
        return data

    # Première ligne = en-têtes
    headers = rows[0]
    data['headers'] = [str(h) if h is not None else f'col_{i}' for i, h in enumerate(headers)]
    data['translated_headers'] = [translate_header(h) for h in data['headers']]

    # Reste = données
    for row in rows[1:]:
        # Créer un dict avec headers traduits
        row_dict = {}
        for i, value in enumerate(row):
            if i < len(data['translated_headers']):
                header = data['translated_headers'][i]
                # Convertir les valeurs
                if value is None:
                    row_dict[header] = None
                elif isinstance(value, (int, float)):
                    row_dict[header] = value
                else:
                    row_dict[header] = str(value)

        # Ajouter seulement si au moins une valeur non-null
        if any(v is not None for v in row_dict.values()):
            data['rows'].append(row_dict)

    return data

def convert_xlsx_file(input_path, output_dir):
    """Convertit un fichier Excel en JSON traduit"""
    try:
        print(f"\n[*] Traitement: {input_path.name}")

        # Charger le workbook
        wb = openpyxl.load_workbook(input_path, data_only=True)

        # Extraire données de toutes les feuilles
        file_data = {
            'source_file': input_path.name,
            'conversion_date': datetime.now().isoformat(),
            'sheets': []
        }

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            print(f"   [*] Feuille: {sheet_name}")

            sheet_data = extract_sheet_data(sheet)

            if sheet_data['rows']:
                file_data['sheets'].append(sheet_data)
                print(f"       [OK] {len(sheet_data['rows'])} lignes extraites")
            else:
                print(f"       [SKIP] Feuille vide")

        # Créer nom de fichier de sortie
        original_name = input_path.stem
        output_name = f"{original_name}_translated.json"
        output_path = output_dir / output_name

        # Sauvegarder JSON
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(file_data, f, indent=2, ensure_ascii=False)

        print(f"   [OK] Sauvegardé: {output_path.name}")
        print(f"   [OK] {len(file_data['sheets'])} feuille(s) traitée(s)")

        return True

    except Exception as e:
        print(f"   [ERREUR] {input_path.name}: {e}")
        return False

def find_xlsx_files(xlsx_dir, exclude_dir):
    """Trouve tous les fichiers Excel sauf ceux dans exclude_dir"""
    xlsx_files = []

    for file_path in xlsx_dir.rglob("*.xlsx"):
        # Exclure les fichiers temporaires Excel
        if file_path.name.startswith('~$'):
            continue

        # Exclure le répertoire import
        if exclude_dir in str(file_path):
            continue

        xlsx_files.append(file_path)

    return sorted(xlsx_files)

def main():
    """Fonction principale"""
    print("="*70)
    print("CONVERSION BATCH FICHIERS EXCEL")
    print("="*70)

    # Vérifier répertoire source
    if not XLSX_DIR.exists():
        print(f"\n[ERREUR] Répertoire source introuvable: {XLSX_DIR}")
        return

    # Créer répertoire de sortie
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    print(f"\n[*] Répertoire source: {XLSX_DIR}")
    print(f"[*] Répertoire sortie: {OUTPUT_DIR}")
    print(f"[*] Exclusion: {EXCLUDE_DIR}")

    # Trouver fichiers Excel
    xlsx_files = find_xlsx_files(XLSX_DIR, EXCLUDE_DIR)

    if not xlsx_files:
        print(f"\n[WARN] Aucun fichier Excel trouvé dans {XLSX_DIR}")
        return

    print(f"\n[*] {len(xlsx_files)} fichier(s) Excel trouvé(s)")

    # Traiter chaque fichier
    success_count = 0
    error_count = 0

    for xlsx_file in xlsx_files:
        if convert_xlsx_file(xlsx_file, OUTPUT_DIR):
            success_count += 1
        else:
            error_count += 1

    # Résumé
    print("\n" + "="*70)
    print("RÉSUMÉ")
    print("="*70)
    print(f"\nFichiers traités: {len(xlsx_files)}")
    print(f"  [OK] Succès: {success_count}")
    print(f"  [ERREUR] Échecs: {error_count}")
    print(f"\nFichiers de sortie: {OUTPUT_DIR}")

    # Lister fichiers créés
    output_files = list(OUTPUT_DIR.glob("*.json"))
    if output_files:
        print(f"\n[*] Fichiers JSON créés ({len(output_files)}):")
        for f in sorted(output_files):
            size_kb = f.stat().st_size / 1024
            print(f"   - {f.name} ({size_kb:.1f} KB)")

    print("\n[OK] Conversion terminée!")

if __name__ == '__main__':
    main()
