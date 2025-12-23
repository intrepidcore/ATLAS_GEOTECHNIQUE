"""
Import des fichiers atlas_import canonisés directement dans la base PostgreSQL
Utilise les sondages existants et crée les échantillons + essais manquants

Usage:
    python import_canonized_xlsx.py <xlsx_file> [--dry-run]
    python import_canonized_xlsx.py --all [--dry-run]

Auteur: Atlas Géotechnique
Version: 1.0.0
"""

import openpyxl
import psycopg2
from psycopg2.extras import execute_values
from pathlib import Path
import argparse
from datetime import datetime
from typing import Dict, List, Optional, Tuple
import uuid

# ============================================================================
# CONFIGURATION
# ============================================================================

DB_CONFIG = {
    'host': 'localhost',
    'port': 5432,
    'database': 'atlas_clean',
    'user': 'atlas',
    'password': 'atlas'
}

RAW_DIR = Path(__file__).parent.parent / 'data' / 'xlsx' / 'RAW'

# ============================================================================
# FONCTIONS UTILITAIRES
# ============================================================================

def normalize_code(code: str) -> str:
    """Normalise un code sondage pour la recherche"""
    if not code:
        return ''
    # Supprimer les suffixes -S1, _S1
    normalized = code.upper().replace('-S1', '').replace('_S1', '').strip()
    # Remplacer les tirets par underscores
    normalized = normalized.replace('-', '_')
    # Supprimer "GRANULO" qui est un artefact du script de canonisation
    if normalized == 'GRANULO':
        return ''
    return normalized


# Mapping manuel des codes canonisés vers les codes en base
CODE_MAPPING = {
    # ADANDOGOU
    'ASSAHOUN-S1': 'ASSAHOUN',
    'BADJA-S1': 'BADJA',
    'KEVE-S1': 'KEVE',
    'KV-S1': 'KEVE',
    # NGOAPO
    'YADE-S1': 'YADE',
    'TCHITCHAO-S1': 'Tchitchao',
    'TCHETCHAO-S1': 'Tchitchao',
    # SOGLO
    'KONSOGOUT1-S1': 'KONSOGOUT1',
    'KONSOGOUT2-S1': 'KONSOGOUT2',
    'NASSABLE-S1': 'NASSABLE',
    'KONTONGBONGUE-S1': 'KONTONGBONGUE',
    # TCHESSI
    'BOHOU-S1': 'BOHOU',
    'LAMA-FEING-S1': 'LAMA_FEING',
    'LAMA-TCHAMD-S1': 'LAMA_TCHAMDE',
    'LAMA-TCHAMDE-S1': 'LAMA_TCHAMDE',
    'LAMA_FEING-S1': 'LAMA_FEING',
    'LAMA_TCHAMDE-S1': 'LAMA_TCHAMDE',
}


def get_connection():
    """Obtient une connexion à la base de données"""
    return psycopg2.connect(**DB_CONFIG)


def find_sondage_id(cursor, code: str, source: str) -> Optional[str]:
    """Trouve l'ID d'un sondage existant par code et source"""
    
    # D'abord, vérifier le mapping manuel
    mapped_code = CODE_MAPPING.get(code.upper(), CODE_MAPPING.get(code, None))
    if mapped_code:
        cursor.execute("""
            SELECT id FROM atlas.sondages 
            WHERE code = %s AND source ILIKE %s
            LIMIT 1
        """, (mapped_code, f'%{source}%'))
        row = cursor.fetchone()
        if row:
            return row[0]
    
    normalized = normalize_code(code)
    if not normalized:
        return None
    
    # Chercher par code exact ou normalisé
    cursor.execute("""
        SELECT id, code FROM atlas.sondages 
        WHERE (UPPER(code) = %s OR UPPER(code) = %s)
        AND source ILIKE %s
        LIMIT 1
    """, (code.upper(), normalized, f'%{source}%'))
    
    row = cursor.fetchone()
    if row:
        return row[0]
    
    # Chercher par code normalisé avec remplacement tirets/underscores
    cursor.execute("""
        SELECT id, code FROM atlas.sondages 
        WHERE UPPER(REPLACE(REPLACE(code, '-', '_'), ' ', '_')) = %s
        AND source ILIKE %s
        LIMIT 1
    """, (normalized, f'%{source}%'))
    
    row = cursor.fetchone()
    return row[0] if row else None


def get_or_create_echantillon(cursor, sondage_id: str, depth_m: float, source: str) -> str:
    """Obtient ou crée un échantillon pour un sondage à une profondeur donnée"""
    
    # Chercher un échantillon existant
    cursor.execute("""
        SELECT id FROM atlas.echantillons 
        WHERE sondage_id = %s AND depth_m = %s
        LIMIT 1
    """, (sondage_id, depth_m))
    
    row = cursor.fetchone()
    if row:
        return row[0]
    
    # Créer un nouvel échantillon
    echantillon_id = str(uuid.uuid4())
    cursor.execute("""
        INSERT INTO atlas.echantillons (id, sondage_id, depth_m, laboratory, created_at)
        VALUES (%s, %s, %s, %s, NOW())
        RETURNING id
    """, (echantillon_id, sondage_id, depth_m, source))
    
    return cursor.fetchone()[0]


# ============================================================================
# PARSERS DE FEUILLES
# ============================================================================

def parse_echantillons_sheet(ws) -> List[Dict]:
    """Parse la feuille echantillons"""
    results = []
    rows = list(ws.iter_rows(values_only=True))
    
    if len(rows) < 2:
        return results
    
    header = [str(h or '').lower().strip() for h in rows[0]]
    
    # Mapper les colonnes
    col_map = {}
    for i, h in enumerate(header):
        if 'code' in h or 'site' in h:
            col_map['code'] = i
        elif 'depth' in h or 'prof' in h:
            col_map['depth'] = i
        elif 'rho' in h or 'densit' in h:
            col_map['rho_s'] = i
        elif 'water' in h or 'teneur' in h:
            col_map['water'] = i
    
    for row in rows[1:]:
        try:
            code = row[col_map.get('code', 0)]
            depth = row[col_map.get('depth', 1)]
            
            if code and depth is not None:
                results.append({
                    'code': str(code),
                    'depth_m': float(depth),
                    'rho_s_gcm3': float(row[col_map['rho_s']]) if col_map.get('rho_s') and row[col_map['rho_s']] else None,
                    'water_content_w': float(row[col_map['water']]) if col_map.get('water') and row[col_map['water']] else None
                })
        except (ValueError, TypeError, IndexError):
            continue
    
    return results


def parse_atterberg_sheet(ws) -> List[Dict]:
    """Parse la feuille atterberg"""
    results = []
    rows = list(ws.iter_rows(values_only=True))
    
    if len(rows) < 2:
        return results
    
    header = [str(h or '').lower().strip() for h in rows[0]]
    
    col_map = {}
    for i, h in enumerate(header):
        if 'code' in h or 'site' in h:
            col_map['code'] = i
        elif 'depth' in h or 'prof' in h:
            col_map['depth'] = i
        elif h == 'wl' or 'liquid' in h:
            col_map['wl'] = i
        elif h == 'wp' or 'plast' in h:
            col_map['wp'] = i
    
    for row in rows[1:]:
        try:
            code = row[col_map.get('code', 0)]
            depth = row[col_map.get('depth', 1)]
            wl = row[col_map.get('wl')] if col_map.get('wl') is not None else None
            wp = row[col_map.get('wp')] if col_map.get('wp') is not None else None
            
            if code and depth is not None and (wl is not None or wp is not None):
                results.append({
                    'code': str(code),
                    'depth_m': float(depth),
                    'wl': float(wl) if wl is not None else None,
                    'wp': float(wp) if wp is not None else None
                })
        except (ValueError, TypeError, IndexError):
            continue
    
    return results


def parse_vbs_sheet(ws) -> List[Dict]:
    """Parse la feuille vbs"""
    results = []
    rows = list(ws.iter_rows(values_only=True))
    
    if len(rows) < 2:
        return results
    
    header = [str(h or '').lower().strip() for h in rows[0]]
    
    col_map = {}
    for i, h in enumerate(header):
        if 'code' in h or 'site' in h:
            col_map['code'] = i
        elif 'depth' in h or 'prof' in h:
            col_map['depth'] = i
        elif 'vbs' in h or 'bleu' in h:
            col_map['vbs'] = i
        elif 'comment' in h:
            col_map['comment'] = i
    
    for row in rows[1:]:
        try:
            code = row[col_map.get('code', 0)]
            depth = row[col_map.get('depth', 1)]
            vbs = row[col_map.get('vbs')] if col_map.get('vbs') is not None else None
            
            if code and depth is not None and vbs is not None:
                results.append({
                    'code': str(code),
                    'depth_m': float(depth),
                    'vbs': float(vbs),
                    'commentaire': str(row[col_map['comment']]) if col_map.get('comment') and row[col_map['comment']] else None
                })
        except (ValueError, TypeError, IndexError):
            continue
    
    return results


def parse_granulo_sheet(ws) -> List[Dict]:
    """Parse la feuille granulo (format long)"""
    results = []
    rows = list(ws.iter_rows(values_only=True))
    
    if len(rows) < 2:
        return results
    
    header = [str(h or '').lower().strip() for h in rows[0]]
    
    col_map = {}
    for i, h in enumerate(header):
        if 'code' in h or 'site' in h:
            col_map['code'] = i
        elif 'depth' in h or 'prof' in h:
            col_map['depth'] = i
        elif 'sieve' in h or 'tamis' in h:
            col_map['sieve'] = i
        elif 'passing' in h or 'passant' in h:
            col_map['passing'] = i
        elif 'method' in h:
            col_map['method'] = i
    
    for row in rows[1:]:
        try:
            code = row[col_map.get('code', 0)]
            depth = row[col_map.get('depth', 1)]
            sieve = row[col_map.get('sieve')] if col_map.get('sieve') is not None else None
            passing = row[col_map.get('passing')] if col_map.get('passing') is not None else None
            
            if code and depth is not None and sieve is not None and passing is not None:
                results.append({
                    'code': str(code),
                    'depth_m': float(depth),
                    'sieve_mm': float(sieve),
                    'passing_pct': float(passing),
                    'method': str(row[col_map['method']]) if col_map.get('method') and row[col_map['method']] else 'tamisage'
                })
        except (ValueError, TypeError, IndexError):
            continue
    
    return results


# ============================================================================
# IMPORT PRINCIPAL
# ============================================================================

def import_xlsx(xlsx_path: Path, dry_run: bool = False) -> Dict:
    """Importe un fichier xlsx canonisé dans la base"""
    
    print(f"\n{'[DRY-RUN] ' if dry_run else ''}Import de: {xlsx_path.name}")
    
    # Extraire la source depuis le nom du fichier
    source_match = xlsx_path.stem.replace('atlas_import_', '')
    
    # Charger le fichier Excel
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    except Exception as e:
        print(f"  ❌ Erreur ouverture fichier: {e}")
        return {'error': str(e)}
    
    # Parser les feuilles
    echantillons_data = []
    atterberg_data = []
    vbs_data = []
    granulo_data = []
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_lower = sheet_name.lower()
        
        if 'echantillon' in sheet_lower:
            echantillons_data = parse_echantillons_sheet(ws)
            print(f"  📋 Feuille echantillons: {len(echantillons_data)} lignes")
        elif 'atterberg' in sheet_lower:
            atterberg_data = parse_atterberg_sheet(ws)
            print(f"  📋 Feuille atterberg: {len(atterberg_data)} lignes")
        elif 'vbs' in sheet_lower:
            vbs_data = parse_vbs_sheet(ws)
            print(f"  📋 Feuille vbs: {len(vbs_data)} lignes")
        elif 'granulo' in sheet_lower:
            granulo_data = parse_granulo_sheet(ws)
            print(f"  📋 Feuille granulo: {len(granulo_data)} lignes")
    
    # Connexion à la base
    conn = get_connection()
    cursor = conn.cursor()
    
    stats = {
        'echantillons_created': 0,
        'atterberg_inserted': 0,
        'vbs_inserted': 0,
        'granulo_inserted': 0,
        'sondages_not_found': set(),
        'errors': []
    }
    
    try:
        # Collecter tous les codes uniques et leurs profondeurs
        all_codes = set()
        code_depths = {}
        
        for item in echantillons_data + atterberg_data + vbs_data:
            code = item['code']
            all_codes.add(code)
            if code not in code_depths:
                code_depths[code] = set()
            code_depths[code].add(item['depth_m'])
        
        for item in granulo_data:
            code = item['code']
            all_codes.add(code)
            if code not in code_depths:
                code_depths[code] = set()
            code_depths[code].add(item['depth_m'])
        
        # Mapper les codes aux sondages existants
        code_to_sondage = {}
        for code in all_codes:
            sondage_id = find_sondage_id(cursor, code, source_match)
            if sondage_id:
                code_to_sondage[code] = sondage_id
            else:
                stats['sondages_not_found'].add(code)
        
        if stats['sondages_not_found']:
            print(f"  ⚠️  Sondages non trouvés: {stats['sondages_not_found']}")
        
        # Créer les échantillons manquants
        echantillon_cache = {}  # (sondage_id, depth_m) -> echantillon_id
        
        for code, depths in code_depths.items():
            if code not in code_to_sondage:
                continue
            
            sondage_id = code_to_sondage[code]
            for depth in depths:
                key = (sondage_id, depth)
                if key not in echantillon_cache:
                    if not dry_run:
                        ech_id = get_or_create_echantillon(cursor, sondage_id, depth, source_match)
                        echantillon_cache[key] = ech_id
                        stats['echantillons_created'] += 1
                    else:
                        echantillon_cache[key] = str(uuid.uuid4())
                        stats['echantillons_created'] += 1
        
        print(f"  ✅ Échantillons créés/trouvés: {stats['echantillons_created']}")
        
        # Insérer les données Atterberg
        for item in atterberg_data:
            code = item['code']
            if code not in code_to_sondage:
                continue
            
            sondage_id = code_to_sondage[code]
            key = (sondage_id, item['depth_m'])
            
            if key not in echantillon_cache:
                continue
            
            echantillon_id = echantillon_cache[key]
            
            if not dry_run:
                try:
                    # Calculer IP
                    ip = None
                    if item['wl'] is not None and item['wp'] is not None:
                        ip = item['wl'] - item['wp']
                    
                    cursor.execute("""
                        INSERT INTO atlas.essais_atterberg (echantillon_id, wl, wp, ip_generated, created_at)
                        VALUES (%s, %s, %s, %s, NOW())
                        ON CONFLICT (echantillon_id) DO UPDATE SET
                            wl = COALESCE(EXCLUDED.wl, atlas.essais_atterberg.wl),
                            wp = COALESCE(EXCLUDED.wp, atlas.essais_atterberg.wp),
                            ip_generated = COALESCE(EXCLUDED.ip_generated, atlas.essais_atterberg.ip_generated)
                    """, (echantillon_id, item['wl'], item['wp'], ip))
                    stats['atterberg_inserted'] += 1
                except Exception as e:
                    stats['errors'].append(f"Atterberg {code}@{item['depth_m']}: {e}")
            else:
                stats['atterberg_inserted'] += 1
        
        print(f"  ✅ Atterberg insérés: {stats['atterberg_inserted']}")
        
        # Insérer les données VBS
        for item in vbs_data:
            code = item['code']
            if code not in code_to_sondage:
                continue
            
            sondage_id = code_to_sondage[code]
            key = (sondage_id, item['depth_m'])
            
            if key not in echantillon_cache:
                continue
            
            echantillon_id = echantillon_cache[key]
            
            if not dry_run:
                try:
                    cursor.execute("""
                        INSERT INTO atlas.essais_vbs (echantillon_id, vbs, commentaire, created_at)
                        VALUES (%s, %s, %s, NOW())
                        ON CONFLICT (echantillon_id) DO UPDATE SET
                            vbs = COALESCE(EXCLUDED.vbs, atlas.essais_vbs.vbs),
                            commentaire = COALESCE(EXCLUDED.commentaire, atlas.essais_vbs.commentaire)
                    """, (echantillon_id, item['vbs'], item.get('commentaire')))
                    stats['vbs_inserted'] += 1
                except Exception as e:
                    stats['errors'].append(f"VBS {code}@{item['depth_m']}: {e}")
            else:
                stats['vbs_inserted'] += 1
        
        print(f"  ✅ VBS insérés: {stats['vbs_inserted']}")
        
        # Insérer les données granulo
        granulo_batch = []
        for item in granulo_data:
            code = item['code']
            if code not in code_to_sondage:
                continue
            
            sondage_id = code_to_sondage[code]
            key = (sondage_id, item['depth_m'])
            
            if key not in echantillon_cache:
                continue
            
            echantillon_id = echantillon_cache[key]
            granulo_batch.append((
                echantillon_id,
                item['sieve_mm'],
                item['passing_pct'],
                item['method']
            ))
        
        if granulo_batch and not dry_run:
            try:
                execute_values(cursor, """
                    INSERT INTO atlas.granulo_points (echantillon_id, sieve_mm, passing_pct, method, created_at)
                    VALUES %s
                    ON CONFLICT (echantillon_id, method, sieve_mm) DO UPDATE SET
                        passing_pct = EXCLUDED.passing_pct
                """, granulo_batch, template="(%s, %s, %s, %s, NOW())")
                stats['granulo_inserted'] = len(granulo_batch)
            except Exception as e:
                stats['errors'].append(f"Granulo batch: {e}")
        else:
            stats['granulo_inserted'] = len(granulo_batch)
        
        print(f"  ✅ Granulo insérés: {stats['granulo_inserted']}")
        
        if not dry_run:
            conn.commit()
            print(f"  💾 Commit effectué")
        else:
            conn.rollback()
            print(f"  🔄 Rollback (dry-run)")
        
        if stats['errors']:
            print(f"  ⚠️  Erreurs: {len(stats['errors'])}")
            for err in stats['errors'][:5]:
                print(f"      - {err}")
    
    except Exception as e:
        conn.rollback()
        print(f"  ❌ Erreur: {e}")
        stats['errors'].append(str(e))
    
    finally:
        cursor.close()
        conn.close()
    
    return stats


def import_all(dry_run: bool = False):
    """Importe tous les fichiers atlas_import_*.xlsx du dossier RAW"""
    
    xlsx_files = list(RAW_DIR.glob('atlas_import_*.xlsx'))
    # Exclure les fichiers du dossier LEGACY
    xlsx_files = [f for f in xlsx_files if 'LEGACY' not in str(f)]
    
    print(f"\n{'[DRY-RUN] ' if dry_run else ''}Import de {len(xlsx_files)} fichiers")
    
    total_stats = {
        'files': 0,
        'echantillons': 0,
        'atterberg': 0,
        'vbs': 0,
        'granulo': 0
    }
    
    for xlsx_file in xlsx_files:
        stats = import_xlsx(xlsx_file, dry_run)
        if 'error' not in stats:
            total_stats['files'] += 1
            total_stats['echantillons'] += stats.get('echantillons_created', 0)
            total_stats['atterberg'] += stats.get('atterberg_inserted', 0)
            total_stats['vbs'] += stats.get('vbs_inserted', 0)
            total_stats['granulo'] += stats.get('granulo_inserted', 0)
    
    print(f"\n{'[DRY-RUN] ' if dry_run else ''}=== RÉSUMÉ ===")
    print(f"  Fichiers traités: {total_stats['files']}")
    print(f"  Échantillons: {total_stats['echantillons']}")
    print(f"  Atterberg: {total_stats['atterberg']}")
    print(f"  VBS: {total_stats['vbs']}")
    print(f"  Granulo: {total_stats['granulo']}")


# ============================================================================
# CLI
# ============================================================================

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Import des fichiers atlas_import canonisés')
    parser.add_argument('xlsx_file', nargs='?', help='Fichier xlsx à importer')
    parser.add_argument('--all', action='store_true', help='Importer tous les fichiers du dossier RAW')
    parser.add_argument('--dry-run', action='store_true', help='Simulation sans écriture')
    
    args = parser.parse_args()
    
    if args.all:
        import_all(args.dry_run)
    elif args.xlsx_file:
        import_xlsx(Path(args.xlsx_file), args.dry_run)
    else:
        parser.print_help()
