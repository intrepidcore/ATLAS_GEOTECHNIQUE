from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, Optional

import openpyxl


RAW_DIR = Path("data/xlsx/amessefe_raw")
OUTPUT_DIR = Path("data/xlsx_convert_md")
OUTPUT_MD = OUTPUT_DIR / "amessefe_raw_compiled.md"


def _anchor(s: str) -> str:
    s = s.strip().lower()
    s = re.sub(r"[^a-z0-9\s_-]", "", s)
    s = re.sub(r"\s+", "-", s)
    s = re.sub(r"-+", "-", s)
    return s or "section"


def _escape_md_cell(v: object) -> str:
    if v is None:
        return ""
    s = str(v)
    # Escape pipes to preserve table structure
    s = s.replace("|", "\\|")
    # Normalize newlines/spaces
    s = s.replace("\r\n", "\n").replace("\r", "\n")
    s = s.replace("\n", "<br/>")
    return s


def _is_row_empty(row: Iterable[object]) -> bool:
    for c in row:
        if c is None:
            continue
        if str(c).strip() != "":
            return False
    return True


def _trim_trailing_empty_columns(rows: list[list[object]]) -> list[list[object]]:
    if not rows:
        return rows

    max_len = max((len(r) for r in rows), default=0)
    if max_len == 0:
        return rows

    # Find last non-empty column index across all rows
    last_non_empty = -1
    for r in rows:
        for i in range(len(r) - 1, -1, -1):
            v = r[i]
            if v is None or str(v).strip() == "":
                continue
            last_non_empty = max(last_non_empty, i)
            break

    if last_non_empty < 0:
        return rows

    new_len = last_non_empty + 1
    return [r[:new_len] + ([""] * max(0, new_len - len(r))) for r in rows]


def _table_md(rows: list[list[object]]) -> str:
    if not rows:
        return "*(feuille vide)*"

    rows = _trim_trailing_empty_columns(rows)
    header = rows[0]
    body = rows[1:]

    # Ensure header not empty; if it is, synthesize column names
    if _is_row_empty(header):
        header = [f"col_{i+1}" for i in range(len(header))]

    header_cells = [_escape_md_cell(v) for v in header]
    sep = ["---" for _ in header_cells]

    out: list[str] = []
    out.append("| " + " | ".join(header_cells) + " |")
    out.append("| " + " | ".join(sep) + " |")

    for r in body:
        # Pad/trim to header length
        if len(r) < len(header_cells):
            r = r + ([""] * (len(header_cells) - len(r)))
        elif len(r) > len(header_cells):
            r = r[: len(header_cells)]

        out.append("| " + " | ".join(_escape_md_cell(v) for v in r) + " |")

    return "\n".join(out)


@dataclass
class SheetDump:
    file_name: str
    sheet_name: str
    rows: list[list[object]]


def _read_sheet_values(ws) -> list[list[object]]:
    rows: list[list[object]] = []
    for row in ws.iter_rows(values_only=True):
        r = list(row)
        if _is_row_empty(r):
            continue
        rows.append(r)
    return rows


def _collect_excels(raw_dir: Path) -> list[Path]:
    files = sorted([p for p in raw_dir.glob("*.xlsx") if not p.name.startswith("~$")])
    return files


def build_compiled_markdown(
    raw_dir: Path,
    files: list[Path],
) -> str:
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    dumps: list[SheetDump] = []
    for fp in files:
        wb = openpyxl.load_workbook(fp, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = _read_sheet_values(ws)
            if not rows:
                continue
            dumps.append(SheetDump(file_name=fp.name, sheet_name=sheet_name, rows=rows))

    lines: list[str] = []
    lines.append("# AMESSEFE Raw Data (compiled)")
    lines.append("")
    lines.append(f"**Source folder:** `{raw_dir.as_posix()}`")
    lines.append(f"**Generated at:** {now}")
    lines.append(f"**Excel files:** {len(files)}")
    lines.append("")
    lines.append("---")
    lines.append("")

    # TOC
    lines.append("## Table des matières")
    lines.append("")

    # Group dumps by file
    file_to_sheets: dict[str, list[SheetDump]] = {}
    for d in dumps:
        file_to_sheets.setdefault(d.file_name, []).append(d)

    for file_name in sorted(file_to_sheets.keys()):
        file_anchor = _anchor(file_name)
        lines.append(f"- [{file_name}](#{file_anchor})")
        for sd in file_to_sheets[file_name]:
            sheet_anchor = _anchor(f"{file_name}-{sd.sheet_name}")
            lines.append(f"  - [{sd.sheet_name}](#{sheet_anchor})")

    lines.append("")
    lines.append("---")
    lines.append("")

    # Content
    for file_name in sorted(file_to_sheets.keys()):
        file_anchor = _anchor(file_name)
        lines.append(f"## {file_name} {{#{file_anchor}}}")
        lines.append("")

        for sd in file_to_sheets[file_name]:
            sheet_anchor = _anchor(f"{file_name}-{sd.sheet_name}")
            lines.append(f"### {sd.sheet_name} {{#{sheet_anchor}}}")
            lines.append("")
            lines.append(f"**Rows (non-empty):** {len(sd.rows) - 1 if len(sd.rows) > 1 else 0}")
            lines.append("")
            lines.append(_table_md(sd.rows))
            lines.append("")
            lines.append("---")
            lines.append("")

    return "\n".join(lines)


def main() -> int:
    if not RAW_DIR.exists():
        raise SystemExit(f"Raw directory not found: {RAW_DIR}")

    files = _collect_excels(RAW_DIR)
    if not files:
        raise SystemExit(f"No .xlsx files found in: {RAW_DIR}")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    md = build_compiled_markdown(RAW_DIR, files)
    OUTPUT_MD.write_text(md, encoding="utf-8")

    print(f"OK: wrote {OUTPUT_MD} ({OUTPUT_MD.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
