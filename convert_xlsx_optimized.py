"""
Script de conversion batch optimisé pour fichiers Excel
Extrait les données et les sauvegarde en CSV par feuille (format compact)
"""

import openpyxl
from pathlib import Path
import csv
from datetime import datetime
import re

# Configuration
XLSX_DIR = Path("data/xlsx")
OUTPUT_DIR = Path("data/xlsx_convert_csv")
EXCLUDE_DIR = "IMPORT"

# Dictionnaire de traduction FR -> EN
TRANSLATIONS = {
    # Colonnes communes
    'code_site': 'site_code',
    'localite': 'locality',
    'date': 'date',
    'profondeur': 'depth_m',
    'depth_m': 'depth_m',
    'laboratoire': 'laboratory',
    'laboratory': 'laboratory',

    # Atterberg
    'wl': 'liquid_limit',
    'wp': 'plastic_limit',
    'ip': 'plasticity_index',
    'limite_liquidite': 'liquid_limit',
    'limite_plasticite': 'plastic_limit',
    'indice_plasticite': 'plasticity_index',

    # VBS
    'vbs': 'methylene_blue_value',
    'bleu_methylene': 'methylene_blue_value',
    'commentaire': 'comment',
    'caracteristique': 'characteristic',

    # Proctor
    'proctor_type': 'proctor_type',
    'gamma_d_max': 'max_dry_density',
    'w_opt': 'optimum_water_content',
    'densite_seche_max': 'max_dry_density',
    'teneur_eau_opt': 'optimum_water_content',

    # Densité
    'rho_s_gcm3': 'absolute_density',
    'rho_d_app': 'apparent_density',
    'densite_absolue': 'absolute_density',
    'densite_apparente': 'apparent_density',
    'water_content_w': 'water_content',
    'teneur_eau': 'water_content',

    # Granulo
    'sieve_mm': 'sieve_mm',
    'tamis_mm': 'sieve_mm',
    'ouverture_tamis': 'sieve_mm',
    'passant': 'passing_pct',
    'passants': 'passing_pct',
    'passant_pct': 'passing_pct',
    'refus': 'retained_pct',
    'refus_pct': 'retained_pct',
    'masse_refus': 'mass_retained_g',
    'mass_refus_cum_g': 'cumulative_mass_retained_g',
    'refus_cum_pct': 'cumulative_retained_pct',

    # Classification
    'hrb': 'hrb_classification',
    'unified': 'uscs_classification',
    'bm': 'methylene_blue_classification',
    'classification_hrb': 'hrb_classification',
    'classification_uscs': 'uscs_classification',

    # Divers
    'adm1': 'region',
    'adm2': 'prefecture',
    'adm3': 'commune',
    'lat': 'latitude',
    'lon': 'longitude',
    'source': 'source',
    'note': 'note',
    'is_index': 'swelling_index',
    'eg': 'swelling_eg',
    'wi': 'swelling_wi',
}

def normalize_header(header):
    """Normalise un en-tête de colonne"""
    if not header:
        return None

    # Convertir en string et nettoyer
    header = str(header).strip().lower()

    # Supprimer caractères spéciaux
    header = re.sub(r'[^\w\s]', '', header)

    # Remplacer espaces par underscore
    header = re.sub(r'\s+', '_', header)

    return header

def translate_header(header):
    """Traduit un en-tête FR -> EN"""
    normalized = normalize_header(header)
    if not normalized:
        return header

    # Chercher traduction exacte
    if normalized in TRANSLATIONS:
        return TRANSLATIONS[normalized]

    # Chercher traduction partielle
    for fr, en in TRANSLATIONS.items():
        if fr in normalized:
            return en

    # Retourner normalisé si pas de traduction
    return normalized

def export_sheet_to_csv(sheet, output_path):
    """Exporte une feuille Excel vers CSV"""
    rows = list(sheet.iter_rows(values_only=True))

    if not rows or len(rows) < 2:
        return 0

    # Première ligne = en-têtes
    headers = rows[0]
    translated_headers = [translate_header(h) if h is not None else f'col_{i}'
                         for i, h in enumerate(headers)]

    # Compter lignes non-vides
    data_rows = []
    for row in rows[1:]:
        # Filtrer lignes complètement vides
        if any(cell is not None and str(cell).strip() for cell in row):
            data_rows.append(row)

    if not data_rows:
        return 0

    # Écrire CSV
    with open(output_path, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(translated_headers)
        writer.writerows(data_rows)

    return len(data_rows)

def convert_xlsx_file(input_path, output_dir):
    """Convertit un fichier Excel en CSV (un par feuille)"""
    try:
        print(f"\n[*] Traitement: {input_path.name}")

        # Charger le workbook
        wb = openpyxl.load_workbook(input_path, data_only=True)

        # Créer sous-dossier pour ce fichier
        file_stem = input_path.stem
        file_output_dir = output_dir / file_stem
        file_output_dir.mkdir(parents=True, exist_ok=True)

        total_rows = 0
        sheets_exported = 0

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]

            # Nettoyer nom de feuille pour nom de fichier
            safe_sheet_name = re.sub(r'[^\w\s-]', '', sheet_name)
            safe_sheet_name = re.sub(r'\s+', '_', safe_sheet_name)

            output_path = file_output_dir / f"{safe_sheet_name}.csv"

            row_count = export_sheet_to_csv(sheet, output_path)

            if row_count > 0:
                total_rows += row_count
                sheets_exported += 1
                print(f"   [OK] {sheet_name}: {row_count} lignes -> {output_path.name}")
            else:
                # Supprimer fichier vide
                output_path.unlink(missing_ok=True)

        # Créer fichier manifest
        manifest_path = file_output_dir / "_manifest.txt"
        with open(manifest_path, 'w', encoding='utf-8') as f:
            f.write(f"Source: {input_path.name}\n")
            f.write(f"Conversion: {datetime.now().isoformat()}\n")
            f.write(f"Sheets: {sheets_exported}\n")
            f.write(f"Total rows: {total_rows}\n")

        print(f"   [OK] {sheets_exported} feuille(s), {total_rows} lignes totales")

        return True

    except Exception as e:
        print(f"   [ERREUR] {input_path.name}: {e}")
        return False

def find_xlsx_files(xlsx_dir, exclude_dir):
    """Trouve tous les fichiers Excel sauf ceux dans exclude_dir"""
    xlsx_files = []

    for file_path in xlsx_dir.rglob("*.xlsx"):
        # Exclure les fichiers temporaires Excel
        if file_path.name.startswith('~$'):
            continue

        # Exclure le répertoire import
        if exclude_dir in str(file_path):
            continue

        xlsx_files.append(file_path)

    return sorted(xlsx_files)

def main():
    """Fonction principale"""
    print("="*70)
    print("CONVERSION OPTIMISÉE FICHIERS EXCEL -> CSV")
    print("="*70)

    # Vérifier répertoire source
    if not XLSX_DIR.exists():
        print(f"\n[ERREUR] Répertoire source introuvable: {XLSX_DIR}")
        return

    # Créer répertoire de sortie
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    print(f"\n[*] Répertoire source: {XLSX_DIR}")
    print(f"[*] Répertoire sortie: {OUTPUT_DIR}")
    print(f"[*] Format: CSV (un fichier par feuille)")
    print(f"[*] Exclusion: {EXCLUDE_DIR}")

    # Trouver fichiers Excel
    xlsx_files = find_xlsx_files(XLSX_DIR, EXCLUDE_DIR)

    if not xlsx_files:
        print(f"\n[WARN] Aucun fichier Excel trouvé dans {XLSX_DIR}")
        return

    print(f"\n[*] {len(xlsx_files)} fichier(s) Excel trouvé(s)")

    # Traiter chaque fichier
    success_count = 0
    error_count = 0

    for xlsx_file in xlsx_files:
        if convert_xlsx_file(xlsx_file, OUTPUT_DIR):
            success_count += 1
        else:
            error_count += 1

    # Résumé
    print("\n" + "="*70)
    print("RÉSUMÉ")
    print("="*70)
    print(f"\nFichiers traités: {len(xlsx_files)}")
    print(f"  [OK] Succès: {success_count}")
    print(f"  [ERREUR] Échecs: {error_count}")
    print(f"\nRépertoire de sortie: {OUTPUT_DIR}")

    # Calculer taille totale
    total_size = 0
    csv_count = 0
    for csv_file in OUTPUT_DIR.rglob("*.csv"):
        total_size += csv_file.stat().st_size
        csv_count += 1

    if csv_count > 0:
        print(f"\n[*] {csv_count} fichiers CSV créés")
        print(f"[*] Taille totale: {total_size / 1024:.1f} KB")
        print(f"[*] Taille moyenne: {total_size / csv_count / 1024:.1f} KB par fichier")

    print("\n[OK] Conversion terminée!")
    print(f"\nChaque fichier Excel a son propre dossier dans {OUTPUT_DIR}")
    print("Chaque feuille est un fichier CSV séparé (format compact)")

if __name__ == '__main__':
    main()
